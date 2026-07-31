#!/usr/bin/env python3
"""
Format a Squarespace orders export (CSV) into a clean, grouped packing list.

What it does, in order:
  1. Fill the Shipping Name DOWN into the blank continuation rows of a
     multi-line order. Squarespace repeats the Order ID on every line item
     but only fills the customer/address fields on the first row; the rest
     are left blank. (If an order has no Shipping Name at all, its Billing
     Name is used so the row still has a label.)
  2. Sort by Shipping Name so every row for the same customer sits together
     -- including a repeat customer's separate orders.
  3. Keep only these columns, in this order:
       Order ID, Lineitem quantity, Lineitem name, Lineitem variant,
       Shipping Name
  4. Write an .xlsx with all text centered, a styled + frozen header,
     auto-sized columns, and light shading that alternates per customer so
     each person's block is easy to see at a glance.

Usage:
    python format_orders.py orders.csv                 # -> orders_formatted.xlsx
    python format_orders.py orders.csv packing.xlsx    # explicit output name
"""

import argparse
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Source column -> header shown in the output, in the exact order requested.
OUTPUT_COLUMNS = [
    ("Order ID", "Order ID"),
    ("Lineitem quantity", "Lineitem quantity"),
    ("Lineitem name", "Lineitem name"),
    ("Lineitem variant", "Lineitem variant"),
    ("Shipping Name", "Shipping Name"),
]

FONT_NAME = "Arial"
HEADER_FILL = "44546A"      # dark slate
HEADER_FONT = "FFFFFF"      # white
BAND_FILL = "F2F2F2"        # very light gray, every other customer block
GRID = "D9D9D9"             # light border color
MAX_COL_WIDTH = 62          # cap so a long item name doesn't blow out the sheet
MIN_COL_WIDTH = 10


def read_rows(csv_path):
    with open(csv_path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        missing = [
            src for src, _ in OUTPUT_COLUMNS if src not in (reader.fieldnames or [])
        ]
        if "Shipping Name" not in (reader.fieldnames or []):
            missing.append("Shipping Name")
        if missing:
            sys.exit(
                "Error: the CSV is missing expected column(s): "
                + ", ".join(sorted(set(missing)))
                + "\nIs this a Squarespace orders export?"
            )
        return list(reader), reader.fieldnames


def fill_shipping_names(rows):
    """Forward-fill Shipping Name across every line of the same Order ID.

    Two passes so it works even if a continuation row somehow precedes the
    row that carries the name: first learn each order's name, then apply it.
    """
    name_by_order = {}
    for r in rows:
        oid = (r.get("Order ID") or "").strip()
        if not oid:
            continue
        ship = (r.get("Shipping Name") or "").strip()
        if ship and oid not in name_by_order:
            name_by_order[oid] = ship
    # Fall back to Billing Name for any order that never had a Shipping Name.
    for r in rows:
        oid = (r.get("Order ID") or "").strip()
        if oid and oid not in name_by_order:
            bill = (r.get("Billing Name") or "").strip()
            if bill:
                name_by_order[oid] = bill
    for r in rows:
        oid = (r.get("Order ID") or "").strip()
        r["Shipping Name"] = name_by_order.get(oid, (r.get("Shipping Name") or "").strip())
    return rows


def as_qty(value):
    """Return an int for a clean whole number, else the original string."""
    s = (value or "").strip()
    try:
        return int(s)
    except (TypeError, ValueError):
        return s


def build_output_rows(rows):
    out = []
    for r in rows:
        out.append(
            {
                "Order ID": (r.get("Order ID") or "").strip(),
                "Lineitem quantity": as_qty(r.get("Lineitem quantity")),
                "Lineitem name": (r.get("Lineitem name") or "").strip(),
                "Lineitem variant": (r.get("Lineitem variant") or "").strip(),
                "Shipping Name": (r.get("Shipping Name") or "").strip(),
            }
        )
    # Group customers together (case-insensitive), then keep each order's
    # line items in their original file order. Python's sort is stable.
    out.sort(key=lambda x: (x["Shipping Name"].casefold(), x["Order ID"]))
    return out


def write_xlsx(out_rows, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"

    headers = [label for _, label in OUTPUT_COLUMNS]
    keys = [src for src, _ in OUTPUT_COLUMNS]

    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color=GRID)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    band_fill = PatternFill("solid", fgColor=BAND_FILL)

    # Header row.
    for c, label in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=label)
        cell.font = Font(name=FONT_NAME, bold=True, color=HEADER_FONT)
        cell.alignment = center
        cell.fill = header_fill
        cell.border = border

    # Data rows, shading alternating per customer block.
    band = False
    prev_name = object()
    for i, row in enumerate(out_rows):
        excel_row = i + 2
        name_key = row["Shipping Name"].casefold()
        if name_key != prev_name:
            band = not band
            prev_name = name_key
        for c, key in enumerate(keys, start=1):
            cell = ws.cell(row=excel_row, column=c, value=row[key])
            cell.font = Font(name=FONT_NAME)
            cell.alignment = center
            cell.border = border
            if band:
                cell.fill = band_fill
            if key == "Order ID":
                cell.number_format = "@"  # keep leading zeros as text

    # Auto-size columns from content (capped).
    for c, (key, label) in enumerate(zip(keys, headers), start=1):
        longest = len(str(label))
        for row in out_rows:
            longest = max(longest, len(str(row[key])))
        width = min(MAX_COL_WIDTH, max(MIN_COL_WIDTH, longest + 2))
        ws.column_dimensions[get_column_letter(c)].width = width

    ws.freeze_panes = "A2"
    if out_rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(out_rows) + 1}"

    wb.save(out_path)


def main():
    ap = argparse.ArgumentParser(
        description=__doc__,
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    ap.add_argument("csv", help="Squarespace orders export (.csv)")
    ap.add_argument(
        "out",
        nargs="?",
        help="Output .xlsx (default: <csv name>_formatted.xlsx)",
    )
    args = ap.parse_args()

    csv_path = Path(args.csv)
    if not csv_path.exists():
        sys.exit(f"Error: no such file: {csv_path}")
    out_path = (
        Path(args.out)
        if args.out
        else csv_path.with_name(csv_path.stem + "_formatted.xlsx")
    )

    rows, _ = read_rows(csv_path)
    rows = fill_shipping_names(rows)
    out_rows = build_output_rows(rows)
    write_xlsx(out_rows, out_path)

    customers = len({r["Shipping Name"].casefold() for r in out_rows})
    print(
        f"Wrote {out_path}  ({len(out_rows)} line items, {customers} customers)"
    )


if __name__ == "__main__":
    main()
